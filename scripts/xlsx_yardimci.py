#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""xlsx_yardimci.py — METİN-güvenli xlsx üretim / okuma / karşılaştırma yardımcısı (Q371②, 2026-10-12).

Proje başına tekrar tekrar yazılan tek-seferlik openpyxl betiklerinin ORTAK çekirdeği. Gate DEĞİL.

İKİ ÖLÇÜLMÜŞ TUZAĞIN KARŞILIĞI:
  ① BAŞTAKİ SIFIRLAR: SAP belge numaraları (10 hane, `0000200008` gibi) sayıya çevrilirse sıfırlar
     SESSİZCE düşer ve anahtar eşleşmesi bozulur. Burada her hücre METİN olarak yazılır
     (`number_format='@'`) ve METİN olarak okunur (pandas `dtype=str` eşdeğeri). Kaynağında zaten
     SAYI olarak saklanmış hücre sayılır ve raporlanır — o hücrenin sıfırı DOSYADA kaybolmuştur,
     okuyucu onu geri getiremez (sessizce "doğru" sayılmaz).
  ② MUTLAK İLİŞKİ HEDEFİ: openpyxl paket ilişkilerini mutlak (`Target="/xl/workbook.xml"`) yazar.
     SAP `cl_fdt_xl` bu biçimi OKUYAMAZ (FDT_EXCEL_CORE005) — `xl/` kökünü Target ile birleştirir.
     `yaz` her üretimden sonra `duzelt` uygular: `_rels/*.rels` içindeki iç hedefler kaynak
     parçaya GÖRELİ yazılır (`xl/worksheets/sheet1.xml` → `worksheets/sheet1.xml`). Harici
     (`TargetMode="External"`) hedeflere ve `[Content_Types].xml` PartName'lerine (OPC gereği
     mutlaktır) DOKUNULMAZ. ⚠ SAP tarafında okunabilirlik bu araçla ÖLÇÜLMEZ (SAP çağrısı yok);
     `denetle` yalnız yapısal ön koşulu (mutlak iç hedef kalmadı) doğrular.

BİLİNÇLİ OLARAK YOK: belge numarası maskeleme (`--redact-pii` benzeri) — 10 haneli SAP belge
numarasını aşırı maskeler, anahtar kıyasını öldürür.

Komutlar (çıkış: 0 tamam · 1 bulgu/fark · 2 kullanım/girdi hatası):
  yaz         --csv <girdi.csv> --cikti <x.xlsx> [--ayrac ;] [--sayfa Ad]
  duzelt      <x.xlsx>                     mutlak iç Target'ları göreliye çevirir
  denetle     <x.xlsx>                     mutlak iç Target var mı (varsa çıkış 1)
  karsilastir <a.xlsx> <b.xlsx> --anahtar <Sütun> [--sayfa Ad] [--csv delta.csv]

`yaz`/`oku`/`karsilastir` openpyxl ister; `duzelt`/`denetle` yalnız standart kütüphaneyle çalışır.
"""
from __future__ import annotations

import argparse
import csv
import datetime as _dt
import os
import posixpath
import re
import sys
import tempfile
import zipfile
from pathlib import Path

if sys.platform == "win32":
    for _a in (sys.stdout, sys.stderr):
        try:
            _a.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

_REL_RE = re.compile(r"<(?:\w+:)?Relationship\b[^>]*>")
_TARGET_RE = re.compile(r'(\bTarget\s*=\s*)(["\'])(.*?)\2')
_DIS_RE = re.compile(r'\bTargetMode\s*=\s*["\']External["\']')


# ── ilişki hedefleri (stdlib) ────────────────────────────────────────────────
def _rels_tabani(rels_adi: str) -> str:
    """`xl/_rels/workbook.xml.rels` → kaynak parçanın klasörü `xl` · `_rels/.rels` → ''."""
    klasor = posixpath.dirname(posixpath.dirname(rels_adi))
    return klasor


def _goreli(target: str, taban: str) -> str:
    hedef = target.lstrip("/")
    return posixpath.relpath(hedef, taban) if taban else hedef


def mutlak_targetlar(yol) -> list[tuple[str, str]]:
    """Paket içindeki MUTLAK iç ilişki hedefleri: [(rels dosyası, Target)]."""
    out = []
    with zipfile.ZipFile(yol) as z:
        for ad in z.namelist():
            if not ad.endswith(".rels"):
                continue
            metin = z.read(ad).decode("utf-8")
            for el in _REL_RE.findall(metin):
                if _DIS_RE.search(el):
                    continue
                m = _TARGET_RE.search(el)
                if m and m.group(3).startswith("/"):
                    out.append((ad, m.group(3)))
    return out


def duzelt(yol) -> int:
    """Mutlak iç Target'ları kaynak parçaya göreli yazar. Döner: düzeltilen hedef sayısı.

    Zip yeniden yazılır (girdi sırası + sıkıştırma türü korunur), önce geçici dosyaya, sonra
    atomik `os.replace` — yarıda kalan yazım asıl dosyayı bozmaz.
    """
    yol = Path(yol)
    sayac = 0
    # mkstemp AÇIK bir tanıtıcı döndürür; kapatılmazsa Windows'ta os.replace/unlink
    # WinError 32 verir (ölçüldü 2026-10-12, bu fixture'ın ilk koşusu).
    fd, _ad = tempfile.mkstemp(suffix=".xlsx", dir=str(yol.parent))
    os.close(fd)
    gecici = Path(_ad)
    try:
        with zipfile.ZipFile(yol) as zin, zipfile.ZipFile(gecici, "w") as zout:
            for bilgi in zin.infolist():
                veri = zin.read(bilgi.filename)
                if bilgi.filename.endswith(".rels"):
                    metin = veri.decode("utf-8")
                    taban = _rels_tabani(bilgi.filename)

                    def _el(m_el):
                        nonlocal sayac
                        el = m_el.group(0)
                        if _DIS_RE.search(el):
                            return el
                        m = _TARGET_RE.search(el)
                        if not m or not m.group(3).startswith("/"):
                            return el
                        sayac += 1
                        return el[:m.start(3)] + _goreli(m.group(3), taban) + el[m.end(3):]

                    veri = _REL_RE.sub(_el, metin).encode("utf-8")
                zout.writestr(bilgi, veri, compress_type=bilgi.compress_type)
        if sayac:
            os.replace(gecici, yol)
        else:
            gecici.unlink()
    except BaseException:
        if gecici.exists():
            gecici.unlink()
        raise
    return sayac


# ── metin-güvenli yazma / okuma (openpyxl) ───────────────────────────────────
def metin_yaz(yol, satirlar, basliklar=None, sayfa="Sayfa1") -> int:
    """Her hücreyi METİN yazar (`@` biçimi), sonra ilişki hedeflerini göreli yapar.
    Döner: düzeltilen Target sayısı (openpyxl sürümüne göre 0 olabilir)."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = sayfa[:31]          # Excel sayfa adı sınırı 31 karakter
    tum = ([list(basliklar)] if basliklar else []) + [list(s) for s in satirlar]
    for r, satir in enumerate(tum, 1):
        for c, deger in enumerate(satir, 1):
            hucre = ws.cell(row=r, column=c, value=None if deger is None else str(deger))
            hucre.number_format = "@"
    wb.save(str(yol))
    return duzelt(yol)


def _metin(deger):
    if deger is None:
        return ""
    if isinstance(deger, bool):
        return "X" if deger else ""
    if isinstance(deger, float) and deger.is_integer():
        return str(int(deger))
    if isinstance(deger, (_dt.datetime, _dt.date, _dt.time)):
        return deger.isoformat()
    return str(deger)


def metin_oku(yol, sayfa=None) -> dict:
    """Tüm hücreleri METİN döndürür + kaynağında SAYI saklanmış hücre sayısı (sıfır kaybı riski)."""
    from openpyxl import load_workbook
    wb = load_workbook(str(yol), read_only=True, data_only=True)
    if not wb.worksheets:
        # ölçüldü: yanlış göreli Target'lı pakette openpyxl hiç sayfa bulamıyor (IndexError yerine)
        raise ValueError("calisma kitabinda okunabilir sayfa yok — iliski hedefleri bozuk olabilir")
    ws = wb[sayfa] if sayfa else wb.worksheets[0]
    satirlar, sayi_hucre = [], 0
    for satir in ws.iter_rows(values_only=True):
        for d in satir:
            if isinstance(d, (int, float)) and not isinstance(d, bool):
                sayi_hucre += 1
        satirlar.append([_metin(d) for d in satir])
    wb.close()
    while satirlar and not any(satirlar[-1]):
        satirlar.pop()
    return {"satirlar": satirlar, "sayi_hucre": sayi_hucre, "sayfa": ws.title}


def karsilastir(a: dict, b: dict, anahtar: str) -> dict:
    """Anahtar sütunlu delta (metin kıyası). Çift anahtar ÖLÇÜLEMEZ — ayrı listelenir."""
    def tablo(t):
        if not t["satirlar"]:
            raise ValueError("bos sayfa")
        bas = t["satirlar"][0]
        if anahtar not in bas:
            raise ValueError("anahtar sutunu yok: %r (basliklar: %s)" % (anahtar, bas))
        i = bas.index(anahtar)
        sozluk, cift = {}, set()
        for s in t["satirlar"][1:]:
            s = s + [""] * (len(bas) - len(s))
            k = s[i]
            if k in sozluk:
                cift.add(k)
            sozluk[k] = dict(zip(bas, s))
        return bas, sozluk, cift

    ba, sa, ca = tablo(a)
    bb, sb, cb = tablo(b)
    cift = sorted(ca | cb)
    ortak_sutun = [c for c in ba if c in bb]
    degisen = []
    for k in sorted(set(sa) & set(sb)):
        if k in cift:
            continue
        fark = [(c, sa[k][c], sb[k][c]) for c in ortak_sutun if sa[k][c] != sb[k][c]]
        if fark:
            degisen.append((k, fark))
    return {"eklenen": sorted(k for k in set(sb) - set(sa) if k not in cift),
            "silinen": sorted(k for k in set(sa) - set(sb) if k not in cift),
            "degisen": degisen, "cift_anahtar": cift,
            "sutun_farki": {"yalniz_a": [c for c in ba if c not in bb], "yalniz_b": [c for c in bb if c not in ba]}}


# ── CLI ──────────────────────────────────────────────────────────────────────
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="metin-guvenli xlsx yardimcisi")
    alt = ap.add_subparsers(dest="komut", required=True)
    y = alt.add_parser("yaz")
    y.add_argument("--csv", required=True)
    y.add_argument("--cikti", required=True)
    y.add_argument("--ayrac", default=";")
    y.add_argument("--sayfa", default="Sayfa1")
    for ad in ("duzelt", "denetle"):
        alt.add_parser(ad).add_argument("xlsx")
    k = alt.add_parser("karsilastir")
    k.add_argument("a")
    k.add_argument("b")
    k.add_argument("--anahtar", required=True)
    k.add_argument("--sayfa")
    k.add_argument("--csv")
    a = ap.parse_args(argv)

    try:
        if a.komut == "yaz":
            with open(a.csv, encoding="utf-8-sig", newline="") as f:
                tum = list(csv.reader(f, delimiter=a.ayrac))
            if not tum:
                print("[KULLANIM] bos CSV: %s" % a.csv)
                return 2
            n = metin_yaz(a.cikti, tum[1:], tum[0], a.sayfa)
            kalan = mutlak_targetlar(a.cikti)
            print("YAZILDI: %s (%d satir, metin) · goreli yapilan Target: %d · kalan mutlak: %d"
                  % (a.cikti, len(tum) - 1, n, len(kalan)))
            return 1 if kalan else 0
        if a.komut == "duzelt":
            n = duzelt(a.xlsx)
            print("DUZELTILDI: %d mutlak ic Target goreliye cevrildi · kalan: %d"
                  % (n, len(mutlak_targetlar(a.xlsx))))
            return 0
        if a.komut == "denetle":
            m = mutlak_targetlar(a.xlsx)
            for ad, t in m:
                print("  [MUTLAK] %s -> %s" % (ad, t))
            print("DENETIM: %d mutlak ic Target (SAP cl_fdt_xl okuyamaz: FDT_EXCEL_CORE005) · "
                  "KAPSAM: yalniz .rels ic hedefleri; SAP'de acilis OLCULMEDI" % len(m))
            return 1 if m else 0
        ta, tb = metin_oku(a.a, a.sayfa), metin_oku(a.b, a.sayfa)
        d = karsilastir(ta, tb, a.anahtar)
    except (OSError, zipfile.BadZipFile, KeyError, ValueError) as e:
        print("[KULLANIM] %s: %s" % (type(e).__name__, e))
        return 2
    except ImportError as e:
        print("[KULLANIM] openpyxl gerekli (%s) — duzelt/denetle openpyxl'siz calisir" % e)
        return 2
    for ad, t in (("A", ta), ("B", tb)):
        if t["sayi_hucre"]:
            print("⚠ %s: %d hucre SAYI olarak saklanmis — bastaki sifir DOSYADA kaybolmus olabilir" % (ad, t["sayi_hucre"]))
    print("KARSILASTIRMA (anahtar=%s): eklenen=%d · silinen=%d · degisen=%d · cift anahtar (OLCULEMEDI)=%d"
          % (a.anahtar, len(d["eklenen"]), len(d["silinen"]), len(d["degisen"]), len(d["cift_anahtar"])))
    if d["sutun_farki"]["yalniz_a"] or d["sutun_farki"]["yalniz_b"]:
        print("  sutun farki: %s" % d["sutun_farki"])
    for k_, fark in d["degisen"][:30]:
        print("  [DEGISEN] %s: %s" % (k_, "; ".join("%s: %r -> %r" % f for f in fark)))
    if a.csv:
        with open(a.csv, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";", lineterminator="\n")
            w.writerow(["durum", a.anahtar, "sutun", "a", "b"])
            for k_ in d["eklenen"]:
                w.writerow(["eklenen", k_, "", "", ""])
            for k_ in d["silinen"]:
                w.writerow(["silinen", k_, "", "", ""])
            for k_, fark in d["degisen"]:
                for c, x, z in fark:
                    w.writerow(["degisen", k_, c, x, z])
            for k_ in d["cift_anahtar"]:
                w.writerow(["cift_anahtar", k_, "", "", ""])
    fark_var = d["eklenen"] or d["silinen"] or d["degisen"] or d["cift_anahtar"]
    return 1 if fark_var else 0


if __name__ == "__main__":
    raise SystemExit(main())
