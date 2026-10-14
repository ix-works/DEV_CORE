#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""xlsx_yardimci — `scripts/xlsx_yardimci.py` korpusu (Q371②, 2026-10-12).

KABUL ÖLÇÜTÜ (kaydın ⚠ RİSK bloğu): `dtype=str` (baştaki sıfırlı 10 haneli belge no round-trip)
+ ilişki Target'ları GÖRELİ. "SAP'de okunabilirlik" ayağı SAP çağrısı ister → bu korpusta
ÖLÇÜLMEZ (ATLANDI satırı basılır); yapısal ön koşulu (mutlak iç hedef kalmadı) ölçülür.

SENARYOLAR (D* yalnız standart kütüphane — CI'da koşar; O* openpyxl ister, yoksa ATLANDI)
  D1  openpyxl-biçimli sahte paket: kök `/xl/workbook.xml` → `xl/workbook.xml`, workbook rels
      `/xl/worksheets/sheet1.xml` → `worksheets/sheet1.xml`, sayfa rels `/xl/drawings/d1.xml` →
      `../drawings/d1.xml`                                                    [--mutasyon-taban]
  D2  Harici (`TargetMode="External"`) hedef DOKUNULMAZ (başı `/` olsa bile)  [--mutasyon-dis]
  D3  `.rels` dışı girdiler bayt-bayt aynı, girdi sırası korunur, [Content_Types] PartName mutlak kalır
  D4  mutlak hedef yokken `duzelt` 0 döner ve dosyayı YENİDEN YAZMAZ (bayt + mtime aynı)
  D5  CLI `denetle`: önce çıkış 1 + [MUTLAK] satırı · `duzelt` sonrası çıkış 0
  O1  metin round-trip: `0000200008`, `0012`, boş, Türkçe → aynen; hücreler metin + `@` biçimi;
      gerçek openpyxl çıktısında mutlak iç hedef KALMADI                        [--mutasyon-sayi]
  O2  kaynağında SAYI saklanmış hücre → metin okunur ama `sayi_hucre` ile RAPORLANIR (sessiz değil)
  O3  anahtarlı kıyas: `0000200008` ≠ `200008` · eklenen/silinen/değişen · çift anahtar ayrı
  B1  3. BAĞLAM — gerçek CLI: `yaz --csv` → `denetle` 0 → `karsilastir` (fark → 1, CSV yazılır)

Koşum:  python tests/fixtures/xlsx_yardimci/run.py [--mutasyon-<kip>]
"""
import importlib.util
import os
import subprocess
import sys
import tempfile
import time
import types
import zipfile
from pathlib import Path

if sys.platform == "win32":
    for _a in (sys.stdout, sys.stderr):
        try:
            _a.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

KOK = Path(__file__).resolve().parents[3]
ARAC = KOK / "scripts" / "xlsx_yardimci.py"

MUTASYON = {
    # rels tabanını yok sayar (her hedefi paket köküne göre yazar) → D1 düşmeli
    "--mutasyon-taban": ('    klasor = posixpath.dirname(posixpath.dirname(rels_adi))\n',
                         '    klasor = ""\n'),
    # Harici hedef korumasını söker → D2 düşmeli
    "--mutasyon-dis": ('                        if _DIS_RE.search(el):\n                            return el\n', ''),
    # rakamlı değeri SAYI yazar (baştaki sıfırı öldürür) → O1 düşmeli
    "--mutasyon-sayi": ('value=None if deger is None else str(deger))',
                        'value=None if deger is None else (int(deger) if str(deger).isdigit() else str(deger)))'),
}
GECERLI_KIP = set(MUTASYON)
_kip = [a for a in sys.argv[1:] if a.startswith("--")]
for a in _kip:
    if a not in GECERLI_KIP:
        print(f"[DURDU] bilinmeyen kip: {a}")
        sys.exit(2)
KIP = _kip[0] if _kip else None


def yukle():
    src = ARAC.read_text(encoding="utf-8")
    if KIP:
        eski, yeni = MUTASYON[KIP]
        if src.count(eski) != 1:
            print(f"[DURDU] mutasyon çapası bulunamadı/çoğul ({KIP}): {eski!r}")
            sys.exit(3)
        src = src.replace(eski, yeni)
    mod = types.ModuleType("xlsx_yardimci_fx")
    mod.__file__ = str(ARAC)
    exec(compile(src, str(ARAC), "exec"), mod.__dict__)
    return mod


X = yukle()
SONUC, ATLANAN = [], []


def kontrol(ad, ok, detay=""):
    SONUC.append((ad, bool(ok), detay))


TMP = Path(tempfile.mkdtemp(prefix="xlsxfx_"))
NS = 'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"'
CT = ('<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
      '<Override PartName="/xl/workbook.xml" ContentType="x"/></Types>')
GIRDILER = [
    ("[Content_Types].xml", CT),
    ("_rels/.rels", f'<Relationships {NS}><Relationship Id="rId1" Type="t" Target="/xl/workbook.xml"/>'
                    f'<Relationship Id="rId2" Type="t" Target="docProps/core.xml"/></Relationships>'),
    ("xl/workbook.xml", "<workbook/>"),
    ("xl/_rels/workbook.xml.rels", f'<Relationships {NS}><Relationship Type="t" Target="/xl/worksheets/sheet1.xml" '
                                   f'Id="rId1" /><Relationship Type="t" Target="styles.xml" Id="rId2" /></Relationships>'),
    ("xl/worksheets/sheet1.xml", "<worksheet>0000200008</worksheet>"),
    ("xl/worksheets/_rels/sheet1.xml.rels",
     f"<Relationships {NS}><Relationship Id='rId1' Type='t' Target='/xl/drawings/d1.xml'/>"
     f'<Relationship Id="rId2" Type="h" Target="/paylasim/dosya.txt" TargetMode="External"/>'
     f'<Relationship Id="rId3" Type="h" Target="https://ornek.invalid/x" TargetMode="External"/></Relationships>'),
    ("xl/drawings/d1.xml", "<d/>"),
]


def sahte_paket(ad, girdiler=GIRDILER):
    p = TMP / ad
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        for n, v in girdiler:
            z.writestr(n, v)
    return p


def oku(p):
    with zipfile.ZipFile(p) as z:
        return [(i.filename, z.read(i.filename).decode("utf-8")) for i in z.infolist()]


# ── D: stdlib ────────────────────────────────────────────────────────────────
p = sahte_paket("d1.xlsx")
once_m = X.mutlak_targetlar(p)
n = X.duzelt(p)
sonra = dict(oku(p))
kontrol("D1 iç hedefler kaynak parçaya GÖRELİ (kök · workbook · sayfa→../)",
        len(once_m) == 3 and n == 3 and X.mutlak_targetlar(p) == []
        and 'Target="xl/workbook.xml"' in sonra["_rels/.rels"]
        and 'Target="worksheets/sheet1.xml"' in sonra["xl/_rels/workbook.xml.rels"]
        and "Target='../drawings/d1.xml'" in sonra["xl/worksheets/_rels/sheet1.xml.rels"],
        f"once={once_m} n={n} sayfa_rels={sonra.get('xl/worksheets/_rels/sheet1.xml.rels')!r}")
sr = sonra["xl/worksheets/_rels/sheet1.xml.rels"]
kontrol("D2 Harici hedef DOKUNULMADI (başı '/' olsa bile)",
        'Target="/paylasim/dosya.txt" TargetMode="External"' in sr and 'Target="https://ornek.invalid/x"' in sr,
        f"{sr!r}")
ilk = dict(GIRDILER)
kontrol("D3 .rels dışı girdiler aynı + sıra korunur + [Content_Types] PartName mutlak",
        [n_ for n_, _ in oku(p)] == [n_ for n_, _ in GIRDILER]
        and all(sonra[k] == v for k, v in ilk.items() if not k.endswith(".rels"))
        and 'PartName="/xl/workbook.xml"' in sonra["[Content_Types].xml"], "")
temiz = sahte_paket("d4.xlsx", [(k, v.replace('Target="/xl/', 'Target="xl/').replace("Target='/xl/", "Target='xl/")
                                  .replace('Target="/xl/worksheets', 'Target="worksheets'))
                                 for k, v in GIRDILER if not k.endswith("sheet1.xml.rels")])
b0, m0 = temiz.read_bytes(), temiz.stat().st_mtime_ns
time.sleep(0.05)
n4 = X.duzelt(temiz)
kontrol("D4 mutlak hedef yok → 0 döner, dosya YENİDEN YAZILMAZ",
        n4 == 0 and temiz.read_bytes() == b0 and temiz.stat().st_mtime_ns == m0 and not list(TMP.glob("tmp*.xlsx")),
        f"n={n4}")
if KIP:
    ATLANAN.append(("D5", "mutasyon kipinde (alt süreç gerçek aracı koşar)"))
else:
    def cli(*a):
        r = subprocess.run([sys.executable, str(ARAC), *a], capture_output=True, text=True,
                           encoding="utf-8", errors="replace", timeout=60)
        return r.returncode, r.stdout + r.stderr
    p5 = sahte_paket("d5.xlsx")
    r1, o1 = cli("denetle", str(p5))
    r2, o2 = cli("duzelt", str(p5))
    r3, o3 = cli("denetle", str(p5))
    kontrol("D5 CLI denetle 1 → duzelt → denetle 0",
            r1 == 1 and o1.count("[MUTLAK]") == 3 and "OLCULMEDI" in o1 and r2 == 0 and r3 == 0,
            f"{r1},{r2},{r3} {o1[-160:]!r}")

# ── O: openpyxl ──────────────────────────────────────────────────────────────
if importlib.util.find_spec("openpyxl") is None:
    for ad in ("O1", "O2", "O3", "B1"):
        ATLANAN.append((ad, "openpyxl yok"))
else:
    from openpyxl import Workbook, load_workbook
    DEGER = [["Belge No", "Kalem", "Açıklama", "Boş"],
             ["0000200008", "0010", "Şişli İğne Çğü", ""],
             ["0000200009", "0012", "ı İ ş Ş", None]]
    x1 = TMP / "o1.xlsx"
    beklenen = [[("" if v is None else v) for v in s] for s in DEGER]
    try:
        X.metin_yaz(x1, DEGER[1:], DEGER[0], "Belgeler")
        t = X.metin_oku(x1)
        ws = load_workbook(str(x1))["Belgeler"]
        hucre = [c for r in ws.iter_rows() for c in r if c.value is not None]
        o1 = (t["satirlar"] == beklenen and t["sayi_hucre"] == 0
              and all(c.data_type == "s" and c.number_format == "@" for c in hucre)
              and X.mutlak_targetlar(x1) == [])
        o1d = f"{t['satirlar']} mutlak={X.mutlak_targetlar(x1)}"
    except Exception as e:   # bozuk paket okunamazsa bu bir FAIL'dir, çökme değil
        o1, o1d = False, f"okunamadi: {type(e).__name__}: {e}"
    kontrol("O1 metin round-trip (baştaki sıfır + Türkçe) + hücre metin/'@' + mutlak iç hedef yok",
            o1, o1d)
    x2 = TMP / "o2.xlsx"
    wb = Workbook()
    wb.active.append(["Belge No", "Miktar"])
    wb.active.append([200008, 1.0])
    wb.save(str(x2))
    t2 = X.metin_oku(x2)
    kontrol("O2 kaynağı SAYI olan hücre metin okunur + sayi_hucre RAPORLANIR",
            t2["satirlar"][1] == ["200008", "1"] and t2["sayi_hucre"] == 2, f"{t2}")
    a_ = {"satirlar": [["Belge No", "Durum"], ["0000200008", "A"], ["0000200009", "A"], ["0000200010", "A"],
                       ["0000200011", "A"], ["0000200011", "B"]], "sayi_hucre": 0}
    b_ = {"satirlar": [["Belge No", "Durum"], ["200008", "A"], ["0000200009", "B"], ["0000200010", "A"],
                       ["0000200012", "A"]], "sayi_hucre": 0}
    d = X.karsilastir(a_, b_, "Belge No")
    kontrol("O3 anahtarlı kıyas: sıfırlı ≠ sıfırsız · değişen · çift anahtar ayrı",
            d["eklenen"] == ["0000200012", "200008"] and d["silinen"] == ["0000200008"]
            and d["degisen"] == [("0000200009", [("Durum", "A", "B")])] and d["cift_anahtar"] == ["0000200011"],
            f"{d}")
    if KIP:
        ATLANAN.append(("B1", "mutasyon kipinde (alt süreç gerçek aracı koşar)"))
    else:
        csvp = TMP / "girdi.csv"
        csvp.write_bytes("﻿Belge No;Açıklama\n0000200008;Şık\n0000200009;İkinci\n".encode("utf-8"))
        csvb = TMP / "girdi_b.csv"
        csvb.write_bytes("Belge No;Açıklama\n0000200008;Şık\n0000200010;Üçüncü\n".encode("utf-8"))
        xa, xb, dc = TMP / "a.xlsx", TMP / "b.xlsx", TMP / "delta.csv"
        r1, o1 = cli("yaz", "--csv", str(csvp), "--cikti", str(xa))
        r1b, _ = cli("yaz", "--csv", str(csvb), "--cikti", str(xb))
        r2, o2 = cli("denetle", str(xa))
        r3, o3 = cli("karsilastir", str(xa), str(xb), "--anahtar", "Belge No", "--csv", str(dc))
        ta = X.metin_oku(xa)
        kontrol("B1 CLI yaz→denetle 0→karsilastir 1 (eklenen 1, silinen 1) + delta CSV",
                r1 == 0 and r1b == 0 and r2 == 0 and r3 == 1 and "eklenen=1 · silinen=1" in o3
                and ta["satirlar"][1] == ["0000200008", "Şık"] and dc.exists()
                and b"\r" not in dc.read_bytes(),
                f"{r1},{r1b},{r2},{r3} {o1[-120:]!r} {o3[-200:]!r}")

ATLANAN.append(("SAP", "cl_fdt_xl ile SAP'de okunabilirlik — SAP çağrısı yok (DOĞRULANAMADI; yapısal ön koşul D1/O1'de)"))

# ── rapor ────────────────────────────────────────────────────────────────────
print(f"xlsx_yardimci{(' ' + KIP) if KIP else ''}")
for ad, ok, detay in SONUC:
    print(f"  [{'PASS' if ok else 'FAIL'}] {ad}")
    if not ok:
        print(f"         -> {detay}")
for ad, sebep in ATLANAN:
    print(f"  [ATLANDI] {ad} — {sebep}")
gecen = sum(1 for _, ok, _ in SONUC if ok)
print(f"TOPLAM: {gecen}/{len(SONUC)} PASS · {len(ATLANAN)} ATLANDI (ölçülmedi; PASS sayılmaz)")
sys.exit(0 if gecen == len(SONUC) else 1)
